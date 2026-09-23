import csv
from django.http import HttpResponse
from django.shortcuts import render
from django.db.models import Q
from django.utils import timezone

from accounts.models import User
from profiles.models import Profile
from universities.models import University, Department, Session, Batch
from locations.models import Division, District, Upazila

from .permissions import admin_required


# ============================================
# FILTERING (same logic as member directory)
# ============================================
def _filtered_profiles(request):
    qs = (
        Profile.objects
        .filter(user__status=User.Status.APPROVED)
        .select_related(
            'user', 'university', 'department', 'session', 'batch',
            'division', 'district', 'upazila',
        )
        .order_by('full_name')
    )

    search = request.GET.get('q', '').strip()
    if search:
        qs = qs.filter(
            Q(full_name__icontains=search) |
            Q(student_id__icontains=search) |
            Q(organization__icontains=search) |
            Q(designation__icontains=search)
        )

    for key, field in [
        ('university', 'university_id'),
        ('department', 'department_id'),
        ('session', 'session_id'),
        ('batch', 'batch_id'),
        ('division', 'division_id'),
        ('district', 'district_id'),
        ('upazila', 'upazila_id'),
    ]:
        value = request.GET.get(key, '')
        if value:
            qs = qs.filter(**{field: value})

    occupation = request.GET.get('occupation', '').strip()
    if occupation:
        qs = qs.filter(occupation__icontains=occupation)

    return qs


# ============================================
# EXPORT COLUMNS
# ============================================
EXPORT_FIELDS = [
    ('Full Name', 'full_name'),
    ('Email', lambda p: p.user.email),
    ('Phone', 'phone'),
    ('Gender', lambda p: p.get_gender_display() or ''),
    ('Date of Birth', lambda p: p.date_of_birth.strftime('%Y-%m-%d') if p.date_of_birth else ''),
    ('Blood Group', 'blood_group'),
    ('University', lambda p: p.university.name if p.university else ''),
    ('Department', lambda p: p.department.name if p.department else ''),
    ('Student ID', 'student_id'),
    ('Session', lambda p: p.session.name if p.session else ''),
    ('Batch', lambda p: p.batch.name if p.batch else ''),
    ('Division', lambda p: p.division.name if p.division else ''),
    ('District', lambda p: p.district.name if p.district else ''),
    ('Upazila', lambda p: p.upazila.name if p.upazila else ''),
    ('Occupation', 'occupation'),
    ('Organization', 'organization'),
    ('Designation', 'designation'),
    ('Facebook', 'facebook'),
    ('LinkedIn', 'linkedin'),
    ('Registered At', lambda p: p.created_at.strftime('%Y-%m-%d %H:%M')),
]


def _get_value(profile, field):
    if callable(field):
        return field(profile)
    return getattr(profile, field, '') or ''


# ============================================
# MAIN VIEW
# ============================================
@admin_required
def export_members(request):
    fmt = request.GET.get('format', '').lower()

    if fmt == 'csv':
        return _export_csv(_filtered_profiles(request))

    if fmt == 'xlsx':
        return _export_xlsx(_filtered_profiles(request))

    # No format → show the filter page
    return render(request, 'admin_dashboard/export.html', _build_context(request))


# ============================================
# CSV EXPORT
# ============================================
def _export_csv(qs):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    filename = f'members_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # Prepend BOM so Excel opens UTF-8 correctly
    response.write('\ufeff')

    writer = csv.writer(response)
    writer.writerow([label for label, _ in EXPORT_FIELDS])

    for profile in qs.iterator():
        writer.writerow([_get_value(profile, field) for _, field in EXPORT_FIELDS])

    return response


# ============================================
# XLSX EXPORT
# ============================================
def _export_xlsx(qs):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return HttpResponse(
            'openpyxl is not installed. Run: pip install openpyxl',
            status=500,
            content_type='text/plain',
        )

    wb = Workbook()
    ws = wb.active
    ws.title = 'Members'

    headers = [label for label, _ in EXPORT_FIELDS]
    ws.append(headers)

    # Style header row
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for profile in qs.iterator():
        ws.append([_get_value(profile, field) for _, field in EXPORT_FIELDS])

    # Auto-fit column widths
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(header)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 3, 45)

    # Freeze header row
    ws.freeze_panes = 'A2'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'members_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ============================================
# CONTEXT FOR FILTER PAGE
# ============================================
def _build_context(request):
    university_id = request.GET.get('university', '')
    department_id = request.GET.get('department', '')
    division_id = request.GET.get('division', '')
    district_id = request.GET.get('district', '')

    departments = Department.objects.filter(is_active=True)
    if university_id:
        departments = departments.filter(university_id=university_id)

    batches = Batch.objects.all()
    if department_id:
        batches = batches.filter(department_id=department_id)

    districts = District.objects.all()
    if division_id:
        districts = districts.filter(division_id=division_id)

    upazilas = Upazila.objects.all()
    if district_id:
        upazilas = upazilas.filter(district_id=district_id)

    return {
        'universities': University.objects.filter(is_active=True),
        'departments': departments,
        'sessions': Session.objects.filter(is_active=True),
        'batches': batches,
        'divisions': Division.objects.all(),
        'districts': districts,
        'upazilas': upazilas,
        'selected': {
            'q': request.GET.get('q', ''),
            'university': university_id,
            'department': department_id,
            'session': request.GET.get('session', ''),
            'batch': request.GET.get('batch', ''),
            'division': division_id,
            'district': district_id,
            'upazila': request.GET.get('upazila', ''),
            'occupation': request.GET.get('occupation', ''),
        },
        'filtered_count': _filtered_profiles(request).count(),
    }